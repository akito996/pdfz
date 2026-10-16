"""Test end-to-end di tutti gli strumenti tramite l'API HTTP."""
from __future__ import annotations

import base64
import datetime as dt
import io
import os
import sys
from pathlib import Path

import pymupdf
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
_WORK = Path(__file__).parent / "_work"
# a ogni esecuzione si riparte da una cartella di lavoro vuota: i file delle esecuzioni vecchie
# (centinaia, dentro OneDrive) rallentano la sincronizzazione e possono restare bloccati
__import__("shutil").rmtree(_WORK, ignore_errors=True)
os.environ.setdefault("HDPDF_WORKDIR", str(_WORK))
os.environ.setdefault("HDPDF_DATADIR", str(Path(__file__).parent / "_data"))
# i pacchetti lingue installati sul PC di build (MSI Lingue) non devono cambiare il risultato dei test
os.environ["HDPDF_ONLY_USER_PACKS"] = "1"

from fastapi.testclient import TestClient  # noqa: E402

from hdpdf.server import create_app  # noqa: E402

client = TestClient(create_app())


# ------------------------------------------------------------------ fixture


def make_pdf(pages=3, text="Documento di prova", extra="") -> bytes:
    d = pymupdf.open()
    for i in range(pages):
        p = d.new_page()
        p.insert_text((72, 72), f"Titolo pagina {i + 1}", fontsize=20)
        p.insert_text((72, 110), f"{text} {i + 1}. Email: mario.rossi@example.com", fontsize=11)
        p.insert_text((72, 130), f"IBAN IT60 X054 2811 1010 0000 0123 456 {extra}", fontsize=11)
        p.insert_link({"kind": pymupdf.LINK_URI, "from": pymupdf.Rect(72, 140, 200, 160),
                       "uri": "https://example.com"})
    img = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 200, 100), 0)
    img.clear_with(180)
    d[0].insert_image(pymupdf.Rect(72, 200, 272, 300), pixmap=img)
    d.set_metadata({"title": "Prova", "author": "Test"})
    return d.tobytes()


def open_doc(data: bytes, name="prova.pdf") -> dict:
    r = client.post("/api/open", files=[("files", (name, data, "application/pdf"))])
    assert r.status_code == 200, r.text
    return r.json()[0]


def run(tool, doc=None, docs=None, expect=200, **params):
    body = {"params": params}
    if doc:
        body["doc"] = doc
    if docs:
        body["docs"] = docs
    r = client.post(f"/api/tools/{tool}", json=body)
    assert r.status_code == expect, f"{tool}: {r.status_code} {r.text}"
    return r.json()


def load(doc_id) -> pymupdf.Document:
    r = client.get(f"/api/items/{doc_id}/file")
    assert r.status_code == 200
    return pymupdf.open("pdf", r.content)


def asset(data: bytes, name: str) -> str:
    r = client.post("/api/assets", files={"file": (name, data)})
    assert r.status_code == 200
    return r.json()["id"]


def png_bytes(w=120, h=60, rgb=(200, 30, 30)) -> bytes:
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, w, h), 0)
    pix.set_rect(pix.irect, rgb)
    return pix.tobytes("png")


@pytest.fixture()
def doc():
    return open_doc(make_pdf())["id"]


# ------------------------------------------------------------------ base


def test_capabilities_and_fonts():
    c = client.get("/api/capabilities").json()
    assert "ocr" in c
    fonts = client.get("/api/fonts").json()
    assert any(f["id"] == "helv" for f in fonts)


# ------------------------------------------------------------------ pagine


def test_merge():
    a, b = open_doc(make_pdf(2))["id"], open_doc(make_pdf(3))["id"]
    res = run("merge", docs=[a, b])
    d = load(res["doc"]["id"])
    assert d.page_count == 5
    assert d.get_toc()[0][1] == "prova"


def test_merge_page_order():
    a, b = open_doc(make_pdf(2, text="AAA"))["id"], open_doc(make_pdf(3, text="BBB"))["id"]
    order = [{"doc": 1, "page": 2}, {"doc": 0, "page": 0, "rotate": 90}, {"doc": 1, "page": 0}]
    d = load(run("merge", docs=[a, b], order=order)["doc"]["id"])
    assert d.page_count == 3
    assert "BBB" in d[0].get_text() and "AAA" in d[1].get_text()
    assert d[1].rotation == 90
    assert [t[2] for t in d.get_toc()] == [1, 2]
    run("merge", docs=[a, b], order=[{"doc": 5, "page": 0}], expect=400)


def test_extract_remove(doc):
    assert load(run("extract_pages", doc, pages="1,3")["doc"]["id"]).page_count == 2
    assert load(run("remove_pages", doc, pages="2")["doc"]["id"]).page_count == 2
    run("remove_pages", doc, pages="1-3", expect=400)


def test_organize(doc):
    res = run("organize", doc, order=[{"page": 2}, {"page": 0, "rotate": 90}, {"page": None}])
    d = load(res["doc"]["id"])
    assert d.page_count == 3
    assert "pagina 3" in d[0].get_text()
    assert d[1].rotation == 90
    assert not d[2].get_text().strip()
    res = run("organize", doc, mode="reverse")
    assert "pagina 3" in load(res["doc"]["id"])[0].get_text()


def test_rotate_crop_resize(doc):
    d = load(run("rotate", doc, angle=90, pages="1")["doc"]["id"])
    assert d[0].rotation == 90 and d[1].rotation == 0
    d = load(run("crop", doc, rect={"x": 50, "y": 50, "w": 300, "h": 200}, pages="1")["doc"]["id"])
    assert round(d[0].rect.width) == 300 and round(d[0].rect.height) == 200
    assert "Titolo pagina 1" in d[0].get_text()
    d = load(run("crop", doc, auto=True)["doc"]["id"])
    assert d[0].rect.width < 595
    d = load(run("resize", doc, format="A3")["doc"]["id"])
    assert round(d[0].rect.width) == 842


def test_page_numbers_and_index(doc):
    d = load(run("page_numbers", doc, format="Pagina {n} di {total}")["doc"]["id"])
    assert "Pagina 2 di 3" in d[1].get_text()
    d = load(run("create_index", doc, source="headings")["doc"]["id"])
    assert d.page_count == 4
    assert "Indice" in d[0].get_text()
    assert d[0].get_links()[0]["page"] == 1


# ------------------------------------------------------------------ contenuto


def test_text_edit(doc):
    spans = run("text_spans", doc, page=0)["data"]["spans"]
    target = next(s for s in spans if s["text"].startswith("Titolo"))
    edit = {"page": 0, "bbox": target["bbox"], "origin": target["origin"], "text": "Titolo MODIFICATO",
            "font": "tibo", "size": target["size"], "color": "#0000ff", "original_font": target["font"]}
    d = load(run("edit_text", doc, edits=[edit])["doc"]["id"])
    t = d[0].get_text()
    assert "Titolo MODIFICATO" in t and "Titolo pagina 1" not in t
    assert "Documento di prova 1" in t  # le righe vicine restano
    # font originale
    edit["font"] = "original"
    edit["text"] = "Titolo pagina X"
    d = load(run("edit_text", doc, edits=[edit])["doc"]["id"])
    assert "Titolo pagina X" in d[0].get_text()


def test_add_text_image_remove_images(doc):
    d = load(run("add_text", doc, items=[{"page": 1, "x": 100, "y": 400, "text": "Nuovo testo\nseconda riga",
                                          "font": "cobo", "size": 14, "color": "#ff0000"}])["doc"]["id"])
    assert "seconda riga" in d[1].get_text()
    img = asset(png_bytes(), "logo.png")
    d = load(run("add_image", doc, image=img, page=1, rect={"x": 50, "y": 500, "w": 100, "h": 50})["doc"]["id"])
    assert len(d[1].get_images()) == 1
    res = run("list_images", doc)
    assert len(res["data"]["images"]) == 1
    d = load(run("remove_images", doc)["doc"]["id"])
    assert all(not pg.get_image_rects(x[0]) for pg in d for x in pg.get_images())


def test_attach(doc):
    a = asset(b"contenuto allegato", "nota.txt")
    d = load(run("attach", doc, files=[a], description="Nota")["doc"]["id"])
    assert d.embfile_names() == ["nota.txt"]


def test_watermark_stamp_sign(doc):
    d = load(run("watermark", doc, text="RISERVATO", layout="tiled")["doc"]["id"])
    assert "RISERVATO" in d[0].get_text()
    d = load(run("watermark", doc, type="image", image=asset(png_bytes(), "w.png"))["doc"]["id"])
    assert d[0].get_images()
    d = load(run("stamp", doc, preset="approvato", pages="1-2", angle=-12)["doc"]["id"])
    assert "APPROVATO" in d[0].get_text() and "APPROVATO" in d[1].get_text()
    sig = client.post("/api/assets/dataurl", json={"data": "data:image/png;base64,"
                                                   + base64.b64encode(png_bytes()).decode()}).json()["id"]
    d = load(run("sign_visual", doc, type="draw", image=sig, page=0,
                 rect={"x": 350, "y": 700, "w": 150, "h": 50})["doc"]["id"])
    assert len(d[0].get_images()) == 2
    d = load(run("sign_visual", doc, type="text", text="Mario Rossi", font="tiit", page=0,
                 rect={"x": 350, "y": 700, "w": 150, "h": 50}, caption=True)["doc"]["id"])
    assert "Mario Rossi" in d[0].get_text()


def test_redact(doc):
    found = run("redact_find", doc, presets=["email", "iban"], terms="Titolo")["data"]["matches"]
    assert len(found) >= 9
    d = load(run("redact", doc, presets=["email"], areas=[{"page": 0, "rect": {"x": 60, "y": 190, "w": 230, "h": 120}}])["doc"]["id"])
    t = "".join(pg.get_text() for pg in d)
    assert "mario.rossi@example.com" not in t
    assert "Documento di prova" in t


# ------------------------------------------------------------------ sicurezza


def test_password_roundtrip(doc):
    res = run("add_password", doc, user_password="segreto", owner_password="admin", permissions=["print"])
    enc = res["doc"]
    assert enc["needs_password"]
    run("remove_password", enc["id"], password="errata", expect=400)
    d = load(run("remove_password", enc["id"], password="segreto")["doc"]["id"])
    assert not d.is_encrypted and d.page_count == 3


def _make_pfx(password=b"test") -> bytes:
    from cryptography import x509
    from cryptography.hazmat.primitives import hashes, serialization
    from cryptography.hazmat.primitives.asymmetric import rsa
    from cryptography.hazmat.primitives.serialization import pkcs12
    from cryptography.x509.oid import NameOID

    key = rsa.generate_private_key(public_exponent=65537, key_size=2048)
    name = x509.Name([x509.NameAttribute(NameOID.COMMON_NAME, "Mario Rossi"),
                      x509.NameAttribute(NameOID.ORGANIZATION_NAME, "PDFz Test")])
    now = dt.datetime.now(dt.timezone.utc)
    cert = (x509.CertificateBuilder().subject_name(name).issuer_name(name).public_key(key.public_key())
            .serial_number(x509.random_serial_number()).not_valid_before(now - dt.timedelta(days=1))
            .not_valid_after(now + dt.timedelta(days=365))
            .add_extension(x509.KeyUsage(True, True, False, False, False, False, False, False, False), True)
            .sign(key, hashes.SHA256()))
    return pkcs12.serialize_key_and_certificates(b"test", key, cert, None,
                                                 serialization.BestAvailableEncryption(password))


def test_sign_cert_validate_remove(doc):
    cert = asset(_make_pfx(), "cert.pfx")
    run("sign_cert", doc, certificate=cert, cert_password="errata", expect=400)
    res = run("sign_cert", doc, certificate=cert, cert_password="test", page=0, reason="Approvazione",
              rect={"x": 350, "y": 650, "w": 200, "h": 70})
    signed = res["doc"]["id"]
    v = run("validate", signed)["data"]
    assert v["signatures"] and v["signatures"][0]["intact"] and v["signatures"][0]["valid"]
    assert "Mario Rossi" in v["signatures"][0]["signer"]
    info = run("info", signed)["data"]
    assert info["security"]["signatures"] == 1
    d = load(run("remove_signature", signed)["doc"]["id"])
    assert not [w for pg in d for w in (pg.widgets() or [])]
    # invisibile
    res = run("sign_cert", doc, certificate=cert, cert_password="test", visible=False)
    assert run("validate", res["doc"]["id"])["data"]["signatures"][0]["intact"]


def test_sanitize(doc):
    d0 = pymupdf.open("pdf", make_pdf())
    d0.embfile_add("x.txt", b"x")
    d0.set_xml_metadata("<x:xmpmeta xmlns:x='adobe:ns:meta/'></x:xmpmeta>")
    cat = d0.pdf_catalog()
    js = d0.get_new_xref()
    d0.update_object(js, "<< /S /JavaScript /JS (app.alert(1);) >>")
    d0.xref_set_key(cat, "OpenAction", f"{js} 0 R")
    src = open_doc(d0.tobytes())["id"]
    d = load(run("sanitize", src, javascript=True, embedded_files=True, xmp_metadata=True, metadata=True,
                 links=True, fonts=True)["doc"]["id"])
    assert not d.embfile_names()
    assert not d.get_xml_metadata()
    assert not d.metadata.get("title")
    assert not d[0].get_links()
    assert d.xref_get_key(d.pdf_catalog(), "OpenAction")[0] == "null"


def test_forms_unlock_flatten():
    d0 = pymupdf.open()
    p = d0.new_page()
    w = pymupdf.Widget()
    w.field_type = pymupdf.PDF_WIDGET_TYPE_TEXT
    w.field_name = "nome"
    w.rect = pymupdf.Rect(72, 72, 300, 100)
    w.field_value = "Mario"
    w.field_flags = 1
    p.add_widget(w)
    p.insert_text((72, 200), "Testo pagina", fontsize=12)
    src = open_doc(d0.tobytes())["id"]
    d = load(run("unlock_forms", src)["doc"]["id"])
    assert not list(d[0].widgets())[0].field_flags & 1
    d = load(run("flatten", src, mode="forms")["doc"]["id"])
    assert not list(d[0].widgets() or [])
    assert "Mario" in d[0].get_text()
    d = load(run("flatten", src, mode="full", dpi=72)["doc"]["id"])
    assert not d[0].get_text().strip() and d[0].get_images()


def test_compress(doc):
    big = pymupdf.open()
    p = big.new_page()
    import random

    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 1500, 1500), 0)
    random.seed(1)
    samples = bytearray(random.getrandbits(8) for _ in range(1500 * 1500 * 3))
    pix = pymupdf.Pixmap(pymupdf.csRGB, 1500, 1500, bytes(samples), 0)
    p.insert_image(pymupdf.Rect(0, 0, 200, 200), pixmap=pix)
    src = open_doc(big.tobytes(deflate=True))["id"]
    res = run("compress", src, level="forte")
    assert res["data"]["after"] < res["data"]["before"] * 0.5
    run("compress", src, level="leggera")


def test_compare():
    a = open_doc(make_pdf(2, "Documento di prova"))["id"]
    b = open_doc(make_pdf(2, "Documento finale"))["id"]
    res = run("compare", docs=[a, b], mode="both")
    s = res["data"]["summary"]
    assert s["modified"] >= 2 and "prova" in res["data"]["changes"][0]["old"]
    assert load(res["doc"]["id"]).page_count == 2


def test_properties_toc_info(doc):
    d = load(run("set_properties", doc, title="Nuovo titolo", author="agjin",
                 creationDate="2025-01-02T10:00:00")["doc"]["id"])
    assert d.metadata["title"] == "Nuovo titolo" and d.metadata["creationDate"].startswith("D:20250102")
    props = run("get_properties", doc)["data"]
    assert props["title"] == "Prova"
    d = load(run("set_toc", doc, toc=[{"level": 1, "title": "Inizio", "page": 1},
                                      {"level": 2, "title": "Dettaglio", "page": 2}])["doc"]["id"])
    assert d.get_toc() == [[1, "Inizio", 1], [2, "Dettaglio", 2]]
    info = run("info", doc)["data"]
    assert info["document"]["pages"] == 3 and info["content"]["images"] == 1
    assert info["security"]["external_links"] == 3
    v = run("validate", doc)["data"]
    assert v["ok"], v["issues"]


# ------------------------------------------------------------------ conversioni


@pytest.mark.parametrize("fmt", ["docx", "xlsx", "pptx", "txt", "html", "png", "csv"])
def test_convert_from_pdf(doc, fmt):
    res = run("convert_from_pdf", doc, format=fmt)
    out = res["output"]
    assert out["size"] > 100
    r = client.get(f"/api/items/{out['id']}/file?download=true")
    assert r.status_code == 200
    if fmt == "docx":
        import docx  # python-docx, dipendenza di pdf2docx

        text = "\n".join(p.text for p in docx.Document(io.BytesIO(r.content)).paragraphs)
        assert "Titolo pagina 1" in text


def test_wait_readable_retries_locked_file(tmp_path, monkeypatch):
    """Un file bloccato per qualche istante (Word, OneDrive, antivirus) non deve far fallire il download."""
    from hdpdf import fileutil

    f = tmp_path / "bloccato.docx"
    f.write_bytes(b"PK test")
    real_open = open
    calls = {"n": 0}

    def flaky_open(path, *a, **k):
        if Path(path) == f and calls["n"] < 3:
            calls["n"] += 1
            raise PermissionError(13, "Permission denied", str(path))
        return real_open(path, *a, **k)

    monkeypatch.setattr("builtins.open", flaky_open)
    fileutil.wait_readable(f, timeout=5)
    assert calls["n"] == 3
    # se resta bloccato oltre il tempo massimo, l'errore arriva comunque (niente attese infinite)
    calls["n"] = -1000
    with pytest.raises(PermissionError):
        fileutil.wait_readable(f, timeout=0.3)


def test_convert_docx_word_fails_falls_back(doc, monkeypatch):
    """Se Microsoft Word c'è ma non riesce (o tiene il file bloccato), si usa il motore integrato."""
    from hdpdf.tools import convert

    def broken_word(src, out, fmt):
        Path(out).write_bytes(b"mezzo file")  # Word lascia un file a metà
        raise PermissionError(13, "Permission denied", str(out))

    monkeypatch.setattr(convert, "ms_office_available", lambda: True)
    monkeypatch.setattr(convert, "_word_pdf_to", broken_word)
    out = run("convert_from_pdf", doc, format="docx")["output"]
    r = client.get(f"/api/items/{out['id']}/file?download=true")
    assert r.status_code == 200
    import docx

    text = "\n".join(p.text for p in docx.Document(io.BytesIO(r.content)).paragraphs)
    assert "Titolo pagina 1" in text


def test_update_check_platforms(monkeypatch):
    """L'avviso di nuova versione c'è per Windows, macOS, Linux e Android (ognuno legge solo il suo file); mai per iPhone."""
    from hdpdf import aggiornamenti

    monkeypatch.setattr(aggiornamenti, "UPDATE_BASE_URL", "https://example.invalid/pdfz/")
    monkeypatch.setattr(aggiornamenti, "_cache", {})
    for p in ("ios", "qualsiasi"):
        assert client.get(f"/api/update/check?platform={p}").json() == {"configured": False, "available": False}
    seen = []

    def fake_urlopen(req, timeout=0):
        seen.append(req.full_url)
        raise OSError("rete non disponibile nel test")

    monkeypatch.setattr(aggiornamenti.urllib.request, "urlopen", fake_urlopen)
    for p in ("windows", "macos", "linux", "android"):
        r = client.get(f"/api/update/check?platform={p}").json()
        assert r["configured"] is True and r["available"] is False
    assert seen == [f"https://example.invalid/pdfz/versione-{p}.json" for p in ("windows", "macos", "linux", "android")]


def test_versione_semver():
    """Versioni MAJOR.MINOR.PATCH: le vecchie AAAAMM.NNN sono più vecchie e i numeri degli installer restano più alti."""
    from hdpdf import versione

    assert versione.newer("1.0.1", "1.0.0") and versione.newer("1.10.0", "1.9.9") and versione.newer("2.0.0", "1.99.99")
    assert versione.newer("1.0.0", "202610.015")          # dalle vecchie versioni si aggiorna
    assert not versione.newer("202610.016", "1.0.0")      # un vecchio file versione-*.json non propone il "ritorno"
    assert not versione.newer("1.0.0", "1.0.0") and not versione.newer("1.0.0", "1.0.1")
    assert versione.msi("1.0.0") == "101.0.0" and versione.msi("1.2.3") == "101.2.3"
    assert versione.android_code("1.0.0") == 1001000000 > 202610015
    assert versione.android_code("1.2.3") == 1001002003


def test_convert_to_pdf():
    # immagine
    r = client.post("/api/open", files=[("files", ("foto.png", png_bytes(), "image/png"))])
    assert r.status_code == 200 and r.json()[0]["pages"] == 1
    # testo
    r = client.post("/api/open", files=[("files", ("note.txt", "ciao àèìòù\nriga 2".encode(), "text/plain"))])
    assert r.status_code == 200
    assert "àèìòù" in load(r.json()[0]["id"])[0].get_text()
    # più immagini in un pdf
    ids = [asset(png_bytes(), "a.png"), asset(png_bytes(rgb=(0, 0, 200)), "b.png")]
    res = client.post("/api/tools/images_to_pdf", json={"params": {"files": ids, "page_size": "A4"}}).json()
    assert load(res["doc"]["id"]).page_count == 2


@pytest.mark.skipif(not __import__("shutil").which("soffice"), reason="LibreOffice assente")
def test_office_to_pdf():
    from openpyxl import Workbook

    wb = Workbook()
    wb.active["A1"] = "Fattura"
    wb.active["B1"] = 1234.5
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post("/api/open", files=[("files", ("tabella.xlsx", buf.getvalue(), "application/octet-stream"))])
    assert r.status_code == 200, r.text
    assert "Fattura" in load(r.json()[0]["id"])[0].get_text()


def test_ocr():
    caps = client.get("/api/capabilities").json()
    if not caps["ocr"]:
        pytest.skip("OCR non disponibile")
    src = pymupdf.open()
    p = src.new_page()
    p.insert_text((72, 100), "Fattura numero 12345 del cliente Rossi", fontsize=18)
    pix = p.get_pixmap(dpi=200)
    scan = pymupdf.open()
    sp = scan.new_page(width=p.rect.width, height=p.rect.height)
    sp.insert_image(sp.rect, pixmap=pix)
    doc = open_doc(scan.tobytes())["id"]
    d = load(run("ocr", doc, languages=["ita", "eng"])["doc"]["id"])
    text = d[0].get_text()
    assert "Fattura" in text and "12345" in text
    assert d[0].search_for("Rossi")


def test_detect_headings_ignores_repeated_headers():
    d0 = pymupdf.open()
    for i in range(3):
        p = d0.new_page()
        p.insert_text((40, 40), "Intestazione Azienda", fontsize=22)
        p.insert_text((40, 90), f"Capitolo {i + 1}", fontsize=16)
        for k in range(10):
            p.insert_text((40, 120 + k * 14), "testo del corpo del documento " * 2, fontsize=10)
    src = open_doc(d0.tobytes())["id"]
    entries = run("detect_headings", src)["data"]["entries"]
    assert [e["title"] for e in entries] == ["Capitolo 1", "Capitolo 2", "Capitolo 3"]
    assert all(e["level"] == 1 for e in entries)


def test_rotated_page_visual_coordinates():
    src = open_doc(make_pdf(1))["id"]
    rot = run("rotate", src, angle=90)["doc"]["id"]
    d = load(run("add_text", rot, items=[{"page": 0, "x": 20, "y": 20, "text": "ANGOLO", "size": 12}])["doc"]["id"])
    pg = d[0]
    r = pg.search_for("ANGOLO")[0] * pg.rotation_matrix  # in coordinate visive
    assert r.x0 < 60 and r.y0 < 60, r


def test_text_edit_wrap_and_move(doc):
    spans = run("text_spans", doc, page=0)["data"]["spans"]
    line = next(s for s in spans if s["text"].startswith("Documento di prova"))
    assert "Email" in line["text"]  # righe intere, non span
    x0, y0, x1, y1 = line["bbox"]
    edit = {"page": 0, "bbox": line["bbox"], "origin": line["origin"], "font": "helv", "size": 11,
            "color": "#000000", "original_font": line["font"], "wrap": True,
            "text": "Prima riga del nuovo paragrafo molto lungo che deve andare a capo nella casella",
            "rect": {"x": x0, "y": y0 + 200, "w": 150, "h": 80}}
    d = load(run("edit_text", doc, edits=[edit])["doc"]["id"])
    words = [w for w in d[0].get_text("words") if w[4] in ("Prima", "casella")]
    assert len({round(w[1]) for w in words}) == 2  # su righe diverse
    assert all(w[1] > y0 + 190 for w in words)  # spostato in basso
    assert "mario.rossi@example.com" not in d[0].get_text()


def test_images_tool(doc):
    imgs = run("list_images", doc)["data"]["images"]
    img = asset(png_bytes(200, 100), "logo.png")
    res = run("images", doc, remove=[{"page": imgs[0]["page"], "xref": imgs[0]["xref"], "rect": dict(zip("xywh", [imgs[0]["bbox"][0], imgs[0]["bbox"][1], imgs[0]["bbox"][2] - imgs[0]["bbox"][0], imgs[0]["bbox"][3] - imgs[0]["bbox"][1]]))}],
              items=[{"image": img, "page": 1, "rect": {"x": 50, "y": 50, "w": 150, "h": 75}, "opacity": 0.5},
                     {"image": img, "page": 0, "rect": {"x": 400, "y": 700, "w": 60, "h": 30}, "all_pages": True}])
    d = load(res["doc"]["id"])
    count = lambda pg: len(pg.get_image_info())
    assert count(d[0]) == 1  # originale rimossa, aggiunta quella "su tutte le pagine"
    assert count(d[1]) == 2 and count(d[2]) == 1
    run("images", doc, expect=400)


@pytest.mark.parametrize("effect", ["blur_light", "blur_heavy"])
def test_redact_blur(doc, effect):
    found = run("redact_find", doc, presets=["email"])["data"]["matches"]
    areas = [{"page": m["page"], "rect": {"x": m["rect"][0], "y": m["rect"][1], "w": m["rect"][2] - m["rect"][0],
                                           "h": m["rect"][3] - m["rect"][1]}} for m in found]
    d = load(run("redact", doc, areas=areas, effect=effect)["doc"]["id"])
    assert "mario.rossi@example.com" not in "".join(pg.get_text() for pg in d)
    assert len(d[1].get_image_info()) == 1  # area sfocata inserita come immagine


def test_redact_blur_rotated_page():
    src = open_doc(make_pdf(1))["id"]
    rot = run("rotate", src, angle=90)["doc"]["id"]
    m = run("redact_find", rot, terms="Titolo")["data"]["matches"][0]
    r = {"x": m["rect"][0], "y": m["rect"][1], "w": m["rect"][2] - m["rect"][0], "h": m["rect"][3] - m["rect"][1]}
    d = load(run("redact", rot, areas=[{"page": 0, "rect": r}], effect="blur_light")["doc"]["id"])
    info = d[0].get_image_info()
    vis = [pymupdf.Rect(i["bbox"]) * d[0].rotation_matrix for i in info]
    assert any(abs(v.x0 - r["x"]) < 2 and abs(v.width - r["w"]) < 2 for v in vis), vis


def test_organize_insert_from_other_pdf(doc):
    other = open_doc(make_pdf(2, "Allegato"))["id"]
    res = run("organize", doc, order=[{"page": 0}, {"src": other, "page": 1}, {"page": 2}, {"page": None}])
    d = load(res["doc"]["id"])
    assert d.page_count == 4 and "Allegato 2" in d[1].get_text()
    res = run("organize", doc, order=[{"page": 2}], new_doc=True, name="estratto.pdf")
    assert res["doc"]["name"] == "estratto.pdf"


def test_signature_slots_and_multi_sign(doc, tmp_path, monkeypatch):
    import hdpdf.config as cfg

    monkeypatch.setattr(cfg, "data_dir", lambda: tmp_path)
    data = "data:image/png;base64," + base64.b64encode(png_bytes(300, 100)).decode()
    assert client.put("/api/signatures/2", json={"data": data}).status_code == 200
    assert client.put("/api/signatures/9", json={"data": data}).status_code == 400
    slots = {s["slot"]: s["data"] for s in client.get("/api/signatures").json()}
    assert slots["2"] and not slots["1"]
    client.delete("/api/signatures/2")
    assert not {s["slot"]: s["data"] for s in client.get("/api/signatures").json()}["2"]
    img = asset(png_bytes(300, 100), "firma.png")
    d = load(run("sign_visual", doc, items=[
        {"type": "image", "image": img, "page": 0, "rect": {"x": 300, "y": 600, "w": 150, "h": 50}},
        {"type": "text", "text": "M. Rossi", "font": "tiit", "size": 20, "page": 1, "rect": {"x": 300, "y": 600, "w": 150, "h": 30}, "all_pages": True},
    ])["doc"]["id"])
    assert len(d[0].get_image_info()) == 2 and "M. Rossi" in d[2].get_text()


# ------------------------------------------------------------------ forme e commenti

def test_shapes_roundtrip(doc):
    items = [
        {"kind": "rect", "page": 0, "rect": {"x": 50, "y": 60, "w": 120, "h": 80}, "stroke": "#dd052b", "fill": "#fdc300", "width": 2, "opacity": 0.8},
        {"kind": "ellipse", "page": 0, "rect": {"x": 200, "y": 60, "w": 100, "h": 60}, "stroke": "#1e396c", "width": 3, "dash": True},
        {"kind": "arrow", "page": 0, "rect": {"x": 60, "y": 200, "w": 150, "h": 50}, "dir": "bl-tr", "stroke": "#000000", "width": 2},
        {"kind": "line", "page": 1, "rect": {"x": 60, "y": 300, "w": 150, "h": 0}, "dir": "l-r", "width": 1},
        {"kind": "check", "page": 0, "rect": {"x": 300, "y": 300, "w": 14, "h": 14}, "stroke": "#008000", "width": 2},
        {"kind": "cross", "page": 0, "rect": {"x": 330, "y": 300, "w": 14, "h": 14}, "width": 2},
    ]
    res = run("shapes", doc, items=items)
    d2 = res["doc"]["id"]
    shapes = run("list_shapes", d2)["data"]["shapes"]
    kinds = sorted(s["kind"] for s in shapes)
    assert kinds == sorted(["rect", "ellipse", "arrow", "line", "check", "cross"])
    r = next(s for s in shapes if s["kind"] == "rect")
    assert abs(r["rect"]["x"] - 50) < 0.5 and abs(r["rect"]["w"] - 120) < 0.5 and r["fill"] == "#fdc300" and r["stroke"] == "#dd052b"
    a = next(s for s in shapes if s["kind"] == "arrow")
    assert a["dir"] == "bl-tr"
    e = next(s for s in shapes if s["kind"] == "ellipse")
    assert e["dash"] and e["width"] == 3
    # modifica (sposta il rettangolo) ed elimina la X
    cross = next(s for s in shapes if s["kind"] == "cross")
    upd = {**r, "rect": {"x": 70, "y": 90, "w": 60, "h": 40}, "orig_page": r["page"]}
    d3 = run("shapes", d2, items=[upd], remove=[{"page": cross["page"], "xref": cross["xref"]}])["doc"]["id"]
    shapes3 = run("list_shapes", d3)["data"]["shapes"]
    assert sorted(s["kind"] for s in shapes3) == sorted(["rect", "ellipse", "arrow", "line", "check"])
    r3 = next(s for s in shapes3 if s["kind"] == "rect")
    assert abs(r3["rect"]["x"] - 70) < 0.5 and abs(r3["rect"]["h"] - 40) < 0.5
    run("shapes", d3, expect=400)


def test_shapes_rotated_page():
    d = pymupdf.open()
    p = d.new_page(width=300, height=500)
    p.set_rotation(90)
    did = open_doc(d.tobytes())["id"]
    res = run("shapes", did, items=[{"kind": "rect", "page": 0, "rect": {"x": 20, "y": 30, "w": 100, "h": 40}}])
    s = run("list_shapes", res["doc"]["id"])["data"]["shapes"][0]
    assert abs(s["rect"]["x"] - 20) < 0.5 and abs(s["rect"]["w"] - 100) < 0.5 and abs(s["rect"]["h"] - 40) < 0.5


def test_comments_flow(doc):
    d1 = run("comments", doc, op="note", page=0, point={"x": 100, "y": 100}, text="Controllare la cifra", author="Mario Rossi")["doc"]["id"]
    # evidenzia il testo del titolo: il riquadro copre la prima riga
    d2 = run("comments", d1, op="markup", kind="highlight", page=0, rect={"x": 40, "y": 40, "w": 500, "h": 50}, text="ok", author="Anna")["doc"]["id"]
    run("comments", d1, op="markup", kind="highlight", page=0, rect={"x": 500, "y": 700, "w": 20, "h": 20}, expect=400)
    d3 = run("comments", d2, op="replace", page=0, rect={"x": 40, "y": 40, "w": 500, "h": 50}, text="Nuovo titolo", author="Anna")["doc"]["id"]
    d4 = run("comments", d3, op="insert", page=0, point={"x": 200, "y": 300}, text="parola", author="Anna")["doc"]["id"]
    cs = run("list_comments", d4)["data"]["comments"]
    types = sorted(c["type"] for c in cs)
    assert types == ["Caret", "Highlight", "StrikeOut", "Text"]
    note = next(c for c in cs if c["type"] == "Text")
    assert note["author"] == "Mario Rossi" and note["content"] == "Controllare la cifra" and note["date"]
    hl = next(c for c in cs if c["type"] == "Highlight")
    assert "Documento" in hl["text"] or hl["text"]
    so = next(c for c in cs if c["type"] == "StrikeOut")
    assert so["content"] == "Sostituire con: Nuovo titolo"
    d5 = run("comments", d4, op="reply", xref=note["xref"], text="Fatto", author="Anna")["doc"]["id"]
    d6 = run("comments", d5, op="state", xref=note["xref"], state="Accepted", author="Anna")["doc"]["id"]
    cs = run("list_comments", d6)["data"]["comments"]
    note = next(c for c in cs if c["type"] == "Text" and not c["irt"])
    assert note.get("review") == "Accepted"
    replies = [c for c in cs if c["irt"] == note["xref"]]
    assert len(replies) == 2 and any(c["content"] == "Fatto" for c in replies)
    d7 = run("comments", d6, op="edit", xref=note["xref"], text="Cifra corretta", author="Mario Rossi")["doc"]["id"]
    d8 = run("comments", d7, op="delete", xref=note["xref"])["doc"]["id"]
    cs = run("list_comments", d8)["data"]["comments"]
    assert not any(c["type"] == "Text" for c in cs)  # eliminate anche risposte e stati


def test_page_words_and_word_selection(doc):
    words = run("page_words", doc, page=0)["data"]["words"]
    assert len(words) > 5 and all({"r", "l"} <= set(w) for w in words)
    # selezione "come in Word": le prime tre parole della prima riga, un riquadro per riga
    first = [w for w in words if w["l"] == words[0]["l"]][:3]
    x0 = min(w["r"]["x"] for w in first); x1 = max(w["r"]["x"] + w["r"]["w"] for w in first)
    y0 = min(w["r"]["y"] for w in first); y1 = max(w["r"]["y"] + w["r"]["h"] for w in first)
    d = run("comments", doc, op="markup", kind="highlight", page=0, rects=[{"x": x0, "y": y0, "w": x1 - x0, "h": y1 - y0}])["doc"]["id"]
    hl = next(c for c in run("list_comments", d)["data"]["comments"] if c["type"] == "Highlight")
    assert hl["text"] and hl["content"] == ""
    run("comments", doc, op="markup", kind="highlight", page=0, rects=[], expect=400)


def test_ocr_with_given_words(doc):
    """App Android: le parole riconosciute arrivano già pronte (Tesseract nell'app) e diventano testo invisibile."""
    blank = pymupdf.open()
    blank.new_page(width=300, height=200)
    did = open_doc(blank.tobytes())["id"]
    prep = run("ocr_prepare", did)["data"]
    assert prep["pages"] and prep["pages"][0]["image"]
    d = run("ocr", did, ocr_words={"0": [[20, 30, 120, 50, "Fattura"], [130, 30, 200, 50, "2026"]]})["doc"]["id"]
    text = pymupdf.open(stream=client.get(f"/api/items/{d}/file").content)[0].get_text()
    assert "Fattura" in text and "2026" in text


def test_owner_registry_and_restricted_unlock(tmp_path):
    import pymupdf
    d = pymupdf.open(); d.new_page()
    r = client.post("/api/open", files={"files": ("a.pdf", d.tobytes(), "application/pdf")})
    doc = r.json()[0]
    # sole restrizioni (password proprietario): il documento si apre ma è "restricted"
    r = client.post("/api/tools/add_password", json={"doc": doc["id"], "params": {"owner_password": "own", "permissions": ["print"]}})
    nd = r.json()["doc"]
    assert nd["restricted"] and not nd["needs_password"] and nd["owner"] and not nd["unlocked"]
    assert client.post(f"/api/docs/{nd['id']}/password", json={"password": "x"}).status_code == 400
    u = client.post(f"/api/docs/{nd['id']}/password", json={"password": "own"}).json()
    assert u["owner_unlocked"] and u["unlocked"]
    # lo stesso file riaperto: proprietario riconosciuto; un file protetto altrove no
    raw = client.get(f"/api/items/{nd['id']}/file").content
    again = client.post("/api/open", files={"files": ("copia.pdf", raw, "application/pdf")}).json()[0]
    assert again["owner"] and again["restricted"]
    other = pymupdf.open(); other.new_page()
    ext = other.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="u", owner_pw="o", permissions=4)
    e = client.post("/api/open", files={"files": ("altro.pdf", ext, "application/pdf")}).json()[0]
    assert e["needs_password"] and not e["owner"]
    u = client.post(f"/api/docs/{e['id']}/password", json={"password": "u"}).json()
    assert u["unlocked"] and not u["owner_unlocked"]


def test_no_password_prompt_for_pdfs_without_real_password():
    """PDF "senza password" per l'utente: non va chiesta nessuna password all'apertura."""
    import pymupdf
    d = pymupdf.open(); d.new_page()
    cases = {
        "normale.pdf": (d.tobytes(), False),
        "cifrato-tutti-i-permessi.pdf": (d.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_128, user_pw="", owner_pw="o", permissions=-1), False),
        "proprietario-vuoto.pdf": (d.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_128, user_pw="", owner_pw="", permissions=4), False),
        "restrizioni.pdf": (d.tobytes(encryption=pymupdf.PDF_ENCRYPT_AES_128, user_pw="", owner_pw="o", permissions=4), True),
    }
    for name, (data, expect) in cases.items():
        info = client.post("/api/open", files={"files": (name, data, "application/pdf")}).json()[0]
        assert not info["needs_password"], name
        assert info["restricted"] is expect, name
