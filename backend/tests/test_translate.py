"""Traduzione offline: pacchetti lingua, motore, PDF tradotto.

Il modello di prova (tests/data/toy-it_en.argosmodel, creato con make_toy_model.py) è minuscolo e casuale:
le "traduzioni" non hanno senso, ma passano per lo stesso percorso dei modelli veri (CTranslate2 + SentencePiece).
"""
from __future__ import annotations

import io
import json
import shutil
import sys
import zipfile
from pathlib import Path

import pymupdf
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent))
from test_tools import client, make_pdf, open_doc, run  # noqa: E402

from hdpdf import translate as tr  # noqa: E402

TOY = Path(__file__).parent / "data" / "toy-it_en.argosmodel"


def retarget(data: bytes, a: str, b: str) -> bytes:
    """Stesso modello, metadati di un'altra coppia (per provare il passaggio dall'inglese)."""
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        for n in src.namelist():
            d = src.read(n)
            if n.endswith("metadata.json"):
                m = json.loads(d)
                m.update(from_code=a, to_code=b)
                d = json.dumps(m).encode()
            z.writestr(n.replace("it_en", f"{a}_{b}"), d)
    return out.getvalue()


@pytest.fixture(scope="module", autouse=True)
def packs():
    shutil.rmtree(tr.user_dir(), ignore_errors=True)
    tr.packs(refresh=True)
    yield
    shutil.rmtree(tr.user_dir(), ignore_errors=True)
    tr.packs(refresh=True)


def test_install_and_languages():
    p = tr.install_archive(TOY.read_bytes())
    assert p.key == "it-en" and (p.path / "model" / "model.bin").exists()
    assert not (p.path / "stanza").exists()
    tr.install_archive(retarget(TOY.read_bytes(), "en", "fr"))
    langs = client.get("/api/translate/languages").json()
    assert set(langs["pairs"]) == {"it-en", "en-fr"}
    assert {l["code"] for l in langs["languages"]} == {"it", "en", "fr"}
    assert [x.key for x in tr.route("it", "fr")] == ["it-en", "en-fr"]  # passa dall'inglese
    with pytest.raises(tr.TranslateError):
        tr.route("fr", "it")
    with pytest.raises(tr.TranslateError):
        tr.install_archive(b"non sono uno zip")


def test_engine_and_helpers():
    out = tr.translate_texts(["Il contratto di prova. Questa riga è un test!", "12345", "mario.rossi@example.com"], "it", "fr")
    assert len(out) == 3 and out[1] == "12345" and out[2] == "mario.rossi@example.com"
    assert isinstance(out[0], str) and out[0]
    assert tr.split_sentences("Prima frase. Seconda frase! Terza.") == ["Prima frase.", "Seconda frase!", "Terza."]
    long = "parola " * 200
    assert all(len(s) <= 401 for s in tr.split_sentences(long))
    assert tr.detect("Il presente contratto è stipulato tra le parti per la fornitura dei servizi") == "it"
    assert tr.detect("This agreement is made between the parties and is valid for one year") == "en"
    assert tr.detect("Der Vertrag wird zwischen den Parteien geschlossen und ist gültig") == "de"
    assert tr.detect("这是一个测试文件") == "zh" and tr.detect("これはテストです") == "ja" and tr.detect("Это тестовый документ") == "ru"


def test_translate_pdf_keeps_layout():
    did = open_doc(make_pdf(2, text="Il contratto di prova per la fornitura"))["id"]
    det = client.post("/api/translate/detect", json={"doc": did}).json()
    assert det["code"] == "it"
    res = run("translate", did, dst="en", src="auto")
    assert res["data"]["src"] == "it" and res["data"]["dst"] == "en" and res["data"]["blocks"] > 0
    assert res["doc"]["name"].endswith("(EN).pdf")
    d = pymupdf.open(stream=client.get(f"/api/items/{res['doc']['id']}/file").content, filetype="pdf")
    assert d.page_count == 2
    assert "Il contratto di prova" not in d[0].get_text()   # testo originale sostituito
    run("translate", did, dst="it", src="it", expect=400)     # stessa lingua
    run("translate", did, dst="de", src="it", expect=400)     # pacchetto mancante


def test_remove_pack():
    assert client.post("/api/translate/remove", json={"key": "en-fr"}).status_code == 200
    assert "en-fr" not in client.get("/api/translate/languages").json()["pairs"]
    assert client.post("/api/translate/remove", json={"key": "xx-yy"}).status_code == 400


def test_export_text():
    did = open_doc(make_pdf(2, text="Il contratto di prova per la fornitura"))["id"]
    res = run("translate", did, dst="en", src="it")
    tid = res["data"]["text_id"]
    for fmt in ("txt", "docx"):
        for layout in ("translation", "bilingual"):
            r = run("translation_export", None, text_id=tid, format=fmt, layout=layout)
            out = client.get(f"/api/items/{r['output']['id']}/file").content
            assert r["output"]["name"].endswith("." + fmt) and len(out) > 50
            if fmt == "txt":
                txt = out.decode("utf-8-sig")
                assert "Pagina 1" in txt and "Pagina 2" in txt
                assert ("Il contratto di prova" in txt) == (layout == "bilingual")
            else:
                from docx import Document
                doc = Document(io.BytesIO(out))
                body = "\n".join(p.text for p in doc.paragraphs) + "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
                assert ("Il contratto di prova" in body) == (layout == "bilingual")
    run("translation_export", None, text_id=tid, format="pdf", expect=400)
