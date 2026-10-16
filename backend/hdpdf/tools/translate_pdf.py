"""Traduzione del PDF mantenendo l'impaginazione: ogni blocco di testo viene sostituito dalla sua
traduzione nella stessa posizione (il carattere si riduce se il testo tradotto è più lungo)."""
from __future__ import annotations

import html
import statistics

import pymupdf

from .. import translate as tr
import json

from ..store import store
from .base import Ctx, Result, ToolError, parse_pages, tool


def _block_info(block: dict) -> dict | None:
    # direzione del testo (coordinate della pagina non ruotata): orizzontale, verso l'alto, capovolto, verso il basso
    dirs = {(1, 0): 0, (0, -1): 90, (-1, 0): 180, (0, 1): 270}
    by_dir: dict[int, list] = {}
    for l in block.get("lines", []):
        key = (round(l["dir"][0]), round(l["dir"][1]))
        if key in dirs and abs(l["dir"][0] - key[0]) < 0.02 and abs(l["dir"][1] - key[1]) < 0.02:
            by_dir.setdefault(dirs[key], []).append(l)
    if not by_dir:
        return None   # testo inclinato: resta com'è
    rotate, lines = max(by_dir.items(), key=lambda kv: sum(len(s["text"]) for l in kv[1] for s in l["spans"]))
    texts, sizes, chars = [], [], 0
    style = {"color": 0, "bold": False, "italic": False, "serif": False, "mono": False}
    best = -1
    for ln in lines:
        t = "".join(s["text"] for s in ln["spans"])
        texts.append(t)
        for s in ln["spans"]:
            n = len(s["text"].strip())
            if not n:
                continue
            sizes += [s["size"]] * n
            chars += n
            if n > best:   # stile del gruppo di testo più lungo
                best = n
                f = s["flags"]
                fn = s.get("font", "").lower()
                style = {"color": s["color"], "bold": bool(f & 16) or "bold" in fn, "italic": bool(f & 2) or "italic" in fn,
                         "serif": bool(f & 4), "mono": bool(f & 8)}
    if not chars:
        return None
    # unisce le righe; "paro-\nla" diventa "parola"
    text = ""
    for t in texts:
        t = t.strip()
        if not t:
            continue
        if text.endswith("-") and t[:1].islower():
            text = text[:-1] + t
        else:
            text = f"{text} {t}" if text else t
    bbox = pymupdf.Rect(block["bbox"])
    xs = [(l["bbox"][0], l["bbox"][2]) for l in lines]
    align = "left"
    if len(xs) > 1 and rotate == 0:
        cx = (bbox.x0 + bbox.x1) / 2
        if all(abs((a + b) / 2 - cx) < 3 for a, b in xs) and max(a for a, _ in xs) - min(a for a, _ in xs) > 6:
            align = "center"
        elif all(abs(b - bbox.x1) < 3 for _, b in xs) and max(a for a, _ in xs) - min(a for a, _ in xs) > 6:
            align = "right"
    if rotate:
        bbox = pymupdf.Rect(lines[0]["bbox"])
        for l in lines[1:]:
            bbox |= pymupdf.Rect(l["bbox"])
    return {"bbox": bbox, "text": text, "size": statistics.median(sizes), "align": align, "lines": len(lines),
            "rotate": rotate, **style}


def _css(b: dict, dst: str) -> str:
    family = "monospace" if b["mono"] else ("serif" if b["serif"] else "sans-serif")
    c = b["color"]
    rgb = f"#{c:06x}" if isinstance(c, int) else "#000000"
    direction = "rtl" if dst in ("ar", "he", "fa", "ur") else "ltr"
    align = b["align"] if direction == "ltr" else ("right" if b["align"] == "left" else b["align"])
    return (f"* {{ font-family: {family}; font-size: {b['size']:.2f}pt; color: {rgb}; line-height: 1.12; margin: 0; padding: 0;"
            f" font-weight: {'bold' if b['bold'] else 'normal'}; font-style: {'italic' if b['italic'] else 'normal'};"
            f" text-align: {align}; direction: {direction}; }}")


@tool("translate")
def translate_pdf(ctx: Ctx) -> Result:
    doc = ctx.open()
    dst = (ctx.params.get("dst") or "").strip()
    src = (ctx.params.get("src") or "auto").strip()
    if not dst:
        raise ToolError("Scegli la lingua in cui tradurre.")
    pages = parse_pages(ctx.params.get("pages"), doc.page_count)
    job = ctx.params.get("job") or ""
    prog = tr.PROGRESS.setdefault(job, {}) if job else {}
    prog.update(label="Lettura del testo…", done=0, total=0)

    plan: list[tuple[int, list[dict]]] = []
    all_text = []
    for pno in pages:
        page = doc[pno]
        blocks = []
        for b in page.get_text("dict", flags=pymupdf.TEXTFLAGS_TEXT)["blocks"]:
            if b.get("type") != 0:
                continue
            info = _block_info(b)
            if info and info["text"].strip():
                blocks.append(info)
                all_text.append(info["text"])
        plan.append((pno, blocks))
    if not all_text:
        doc.close()
        raise ToolError("Nel documento non c'è testo da tradurre. Se è una scansione, esegui prima l'OCR.")
    if src == "auto":
        src = tr.detect(" ".join(all_text))
    if src == dst:
        doc.close()
        raise ToolError(f"Il documento sembra già in {tr.LANG_NAMES.get(dst, dst)}: scegli un'altra lingua o indica la lingua d'origine.")
    try:
        tr.route(src, dst)
    except tr.TranslateError as ex:
        doc.close()
        raise ToolError(str(ex))

    prog.update(label=f"Traduzione {tr.LANG_NAMES.get(src, src)} → {tr.LANG_NAMES.get(dst, dst)}…")
    texts = [b["text"] for _, blocks in plan for b in blocks]
    try:
        out = tr.translate_texts(texts, src, dst, progress=prog)
    except tr.TranslateError as ex:
        doc.close()
        raise ToolError(str(ex))

    prog.update(label="Impaginazione…")
    pairs = []
    i = 0
    for pno, blocks in plan:
        page = doc[pno]
        page.clean_contents()
        for b in blocks:
            b["out"] = out[i]
            pairs.append({"page": pno, "src": b["text"], "dst": out[i]})
            i += 1
            # via il testo originale, senza toccare immagini e grafica
            page.add_redact_annot(b["bbox"], fill=False)
        page.apply_redactions(images=pymupdf.PDF_REDACT_IMAGE_NONE, graphics=pymupdf.PDF_REDACT_LINE_ART_NONE)
        for b in blocks:
            if not b["out"].strip():
                continue
            # coordinate della pagina non ruotata (come le dà get_text); il testo segue la direzione originale
            r = pymupdf.Rect(b["bbox"])
            grow = b["size"] * 0.25
            mb = page.rect * page.derotation_matrix   # pagina non ruotata
            mb.normalize()
            if b["rotate"] in (0, 180):
                r.x1 = min(r.x1 + 2, mb.x1)       # un filo di margine per non spezzare l'ultima parola
                r.y1 = min(r.y1 + grow, mb.y1)
            else:
                r.y1 = min(r.y1 + 2, mb.y1)
                r.x1 = min(r.x1 + grow, mb.x1)
            content = html.escape(b["out"]).replace("\n", "<br>")
            try:
                page.insert_htmlbox(r, f"<p>{content}</p>", css=_css(b, dst), scale_low=0, rotate=b["rotate"])
            except Exception:  # noqa: BLE001 - un blocco problematico non deve fermare tutto
                page.insert_textbox(r, b["out"], fontsize=max(4, b["size"] * 0.8), fontname="helv", align=0)
    prog.update(label="Fatto", done=prog.get("total", 0))
    stem = ctx.item.name.rsplit(".", 1)[0]
    name = f"{stem} ({dst.upper()}).pdf"
    # testo completo (originale + traduzione) conservato per l'esportazione in TXT/Word
    text_item = store.add_bytes(json.dumps({"src": src, "dst": dst, "name": stem, "pairs": pairs}, ensure_ascii=False).encode("utf-8"),
                                f"{stem}-traduzione.json", kind="asset")
    data = {"src": src, "dst": dst, "src_name": tr.LANG_NAMES.get(src, src), "dst_name": tr.LANG_NAMES.get(dst, dst),
            "pages": [p + 1 for p in pages], "pairs": pairs[:3000], "blocks": len(pairs), "text_id": text_item.id}
    return Result(pdf=ctx.save(doc), name=name, data=data,
                  message=f"Tradotto da {data['src_name']} a {data['dst_name']} ({len(pairs)} blocchi di testo).")


@tool("translation_export", needs_doc=False)
def translation_export(ctx: Ctx) -> Result:
    """Salva il testo tradotto: params.text_id, format = txt | docx, layout = translation | bilingual."""
    item = ctx.asset("text_id")
    t = json.loads(item.path.read_text(encoding="utf-8"))
    fmt = (ctx.params.get("format") or "txt").lower()
    bilingual = ctx.params.get("layout") == "bilingual"
    src_name, dst_name = tr.LANG_NAMES.get(t["src"], t["src"]), tr.LANG_NAMES.get(t["dst"], t["dst"])
    pairs = t["pairs"]
    suffix = f"{t['name']} ({t['dst'].upper()})" + (" bilingue" if bilingual else "")
    if fmt == "txt":
        lines, page = [], None
        for p in pairs:
            if p["page"] != page:
                page = p["page"]
                lines.append(f"\n--- Pagina {page + 1} ---\n")
            if bilingual:
                lines.append(f"[{src_name}] {p['src']}\n[{dst_name}] {p['dst']}\n")
            else:
                lines.append(p["dst"] + "\n")
        out = ctx.out(".txt")
        out.write_text("\ufeff" + "\n".join(lines).strip() + "\n", encoding="utf-8")   # BOM: Blocco note mostra gli accenti
        return Result(output=out, output_name=f"{suffix}.txt", message="Testo tradotto pronto.")
    if fmt == "docx":
        from docx import Document
        from docx.shared import Pt, RGBColor

        d = Document()
        d.core_properties.title = suffix
        d.add_heading(f"{t['name']} — traduzione {src_name} → {dst_name}", level=1)
        if bilingual:
            table = d.add_table(rows=1, cols=2)
            table.style = "Table Grid"
            table.rows[0].cells[0].text, table.rows[0].cells[1].text = src_name, dst_name
            for c in table.rows[0].cells:
                for r in c.paragraphs[0].runs:
                    r.bold = True
        page = None
        for p in pairs:
            if p["page"] != page:
                page = p["page"]
                if bilingual:
                    row = table.add_row().cells
                    row[0].merge(row[1]).text = f"Pagina {page + 1}"
                    row[0].paragraphs[0].runs[0].italic = True
                else:
                    h = d.add_paragraph(f"Pagina {page + 1}")
                    h.runs[0].italic = True
                    h.runs[0].font.size = Pt(9)
                    h.runs[0].font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
            if bilingual:
                row = table.add_row().cells
                row[0].text, row[1].text = p["src"], p["dst"]
            else:
                d.add_paragraph(p["dst"])
        out = ctx.out(".docx")
        d.save(out)
        return Result(output=out, output_name=f"{suffix}.docx", message="Documento Word con il testo tradotto pronto.")
    raise ToolError("Formato non supportato.")
