"""Conversioni: PDF -> formati Office (Word, Excel, PowerPoint, OpenDocument, testo,
immagini, PDF/A) e altri formati -> PDF.

Motori usati, in ordine di preferenza:
* Microsoft Office via COM (Windows, se installato): massima fedeltà per Office -> PDF
  e per PDF -> Word.
* LibreOffice headless (se installato).
* Librerie Python integrate (pdf2docx, pdfplumber+openpyxl, python-pptx, PyMuPDF),
  sempre disponibili.
"""
from __future__ import annotations

import io

import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path

import pymupdf

from ..config import ghostscript_cmd, ms_office_available, soffice_cmd
from ..fileutil import move_when_free, wait_readable
from ..store import store
from .base import Ctx, Result, ToolError, boolean, num, parse_pages, tool

IMAGE_EXT = {".png", ".jpg", ".jpeg", ".bmp", ".gif", ".tif", ".tiff", ".webp", ".jxr", ".pnm", ".pbm", ".pgm", ".ppm"}
WORD_EXT = {".doc", ".docx", ".docm", ".dot", ".dotx", ".rtf", ".odt", ".wpd", ".xml"}
EXCEL_EXT = {".xls", ".xlsx", ".xlsm", ".xlsb", ".ods", ".csv"}
PPT_EXT = {".ppt", ".pptx", ".pptm", ".pps", ".ppsx", ".odp"}
TEXT_EXT = {".txt", ".md", ".log", ".json", ".ini", ".cfg"}
HTML_EXT = {".html", ".htm", ".xhtml"}
MUPDF_DOC_EXT = {".xps", ".oxps", ".epub", ".fb2", ".cbz", ".mobi", ".svg"}

SUPPORTED_INPUT = sorted(IMAGE_EXT | WORD_EXT | EXCEL_EXT | PPT_EXT | TEXT_EXT | HTML_EXT | MUPDF_DOC_EXT | {".pdf"})


# ------------------------------------------------------------------ motori esterni


def _libreoffice(src: Path, target: str, outdir: Path, filter_opt: str | None = None) -> Path:
    exe = soffice_cmd()
    if not exe:
        raise ToolError("Per questa conversione serve LibreOffice o Microsoft Office installato sul PC.")
    profile = Path(tempfile.mkdtemp(prefix="hdpdf-lo-"))
    conv = target + (f":{filter_opt}" if filter_opt else "")
    args = [exe, f"-env:UserInstallation={profile.as_uri()}", "--headless", "--norestore",
            "--convert-to", conv, "--outdir", str(outdir), str(src)]
    if target in ("docx", "odt", "rtf", "doc") and src.suffix.lower() == ".pdf":
        args.insert(-1, "--infilter=writer_pdf_import")
    try:
        subprocess.run(args, check=True, timeout=300, capture_output=True)
    except subprocess.TimeoutExpired:
        raise ToolError("LibreOffice non ha risposto in tempo.")
    except subprocess.CalledProcessError as ex:
        raise ToolError(f"Conversione LibreOffice non riuscita: {ex.stderr.decode(errors='ignore')[:300]}")
    finally:
        shutil.rmtree(profile, ignore_errors=True)
    out = outdir / (src.stem + "." + target.split(":")[0])
    if not out.exists():
        cands = list(outdir.glob(src.stem + ".*"))
        cands = [c for c in cands if c != src]
        if not cands:
            raise ToolError("LibreOffice non ha prodotto alcun file.")
        out = cands[0]
    return out


def _office_to_pdf(src: Path, out: Path) -> None:
    """Esporta in PDF con Microsoft Office (COM). Solo Windows."""
    import pythoncom  # type: ignore
    import win32com.client  # type: ignore

    ext = src.suffix.lower()
    pythoncom.CoInitialize()
    try:
        if ext in WORD_EXT or ext in TEXT_EXT or ext in HTML_EXT:
            app = win32com.client.DispatchEx("Word.Application")
            app.Visible = False
            app.DisplayAlerts = 0
            try:
                d = app.Documents.Open(str(src), ConfirmConversions=False, ReadOnly=True, AddToRecentFiles=False)
                d.ExportAsFixedFormat(str(out), 17, OptimizeFor=0, CreateBookmarks=1, DocStructureTags=True)
                d.Close(False)
            finally:
                app.Quit()
        elif ext in EXCEL_EXT:
            app = win32com.client.DispatchEx("Excel.Application")
            app.Visible = False
            app.DisplayAlerts = False
            try:
                wb = app.Workbooks.Open(str(src), ReadOnly=True, UpdateLinks=0)
                wb.ExportAsFixedFormat(0, str(out))
                wb.Close(False)
            finally:
                app.Quit()
        elif ext in PPT_EXT:
            app = win32com.client.DispatchEx("PowerPoint.Application")
            try:
                p = app.Presentations.Open(str(src), ReadOnly=True, WithWindow=False)
                p.SaveAs(str(out), 32)
                p.Close()
            finally:
                app.Quit()
        else:
            raise ToolError(f"Formato non gestito da Office: {ext}")
    finally:
        pythoncom.CoUninitialize()


def _word_pdf_to(src: Path, out: Path, fmt: int) -> None:
    """Apre un PDF con Word (PDF Reflow) e lo salva come docx(16)/rtf(6)/odt(23)/doc(0).

    Word salva in una cartella temporanea sua (fuori da OneDrive): se salvasse direttamente in una cartella
    sincronizzata, OneDrive/Office terrebbero il .docx bloccato per qualche secondo e chi lo apre subito dopo
    riceve "PermissionError [Errno 13]". Il file si sposta al suo posto solo quando Word lo ha davvero rilasciato.
    """
    import pythoncom  # type: ignore
    import win32com.client  # type: ignore

    tmpdir = Path(tempfile.mkdtemp(prefix="hdpdf-word-"))
    tmp_out = tmpdir / ("documento" + out.suffix)
    pythoncom.CoInitialize()
    try:
        app = win32com.client.DispatchEx("Word.Application")
        app.Visible = False
        app.DisplayAlerts = 0
        try:
            d = app.Documents.Open(str(src), ConfirmConversions=False, ReadOnly=True, AddToRecentFiles=False)
            d.SaveAs2(str(tmp_out), FileFormat=fmt)
            d.Close(False)
            d = None
        finally:
            app.Quit()
            app = None  # rilascia subito il collegamento COM: Word chiude prima il file
    finally:
        pythoncom.CoUninitialize()
    try:
        if not tmp_out.exists():
            raise ToolError("Microsoft Word non ha prodotto il file.")
        wait_readable(tmp_out)
        move_when_free(tmp_out, out)
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)


# ------------------------------------------------------------------ altri formati -> PDF


def _image_pdf(src: Path, ext: str) -> pymupdf.Document:
    """Immagine -> PDF in memoria. Il JPEG resta com'è (nessuna ricompressione): veloce e senza perdita."""
    img = pymupdf.open(src, filetype=ext[1:])  # gestisce anche TIFF multipagina
    pdf_bytes = img.convert_to_pdf()
    img.close()
    return pymupdf.open("pdf", pdf_bytes)


def any_to_pdf(src: Path, name: str, engine: str = "auto") -> Path:
    """Converte un file qualsiasi supportato in PDF e restituisce il percorso."""
    ext = Path(name).suffix.lower()
    out = store.new_path(".pdf")
    if ext == ".pdf":
        shutil.copy(src, out)
        return out
    if ext in IMAGE_EXT:
        pdf = _image_pdf(src, ext)
        pdf.save(out, garbage=1, deflate=True)   # solo immagini: niente da "ripulire", salvataggio più rapido
        pdf.close()
        return out
    if ext in MUPDF_DOC_EXT:
        d = pymupdf.open(src, filetype=ext[1:])
        pdf = pymupdf.open("pdf", d.convert_to_pdf())
        pdf.save(out, garbage=3, deflate=True)
        return out
    if ext in TEXT_EXT and engine != "office":
        return _text_to_pdf(src.read_text(encoding="utf-8", errors="replace"), out, html=False)
    if ext in HTML_EXT and engine == "builtin":
        return _text_to_pdf(src.read_text(encoding="utf-8", errors="replace"), out, html=True)
    office_like = WORD_EXT | EXCEL_EXT | PPT_EXT | HTML_EXT | TEXT_EXT
    if ext in office_like:
        # il file deve avere l'estensione giusta perché Office/LibreOffice lo riconoscano
        tmpdir = Path(tempfile.mkdtemp(prefix="hdpdf-conv-", dir=store.root))
        named = tmpdir / Path(name).name
        shutil.copy(src, named)
        try:
            if engine in ("auto", "office") and ms_office_available():
                try:
                    _office_to_pdf(named, out)
                    return out
                except ToolError:
                    raise
                except Exception:
                    if engine == "office":
                        raise ToolError("Conversione con Microsoft Office non riuscita.")
            if ext in HTML_EXT and not soffice_cmd():
                return _text_to_pdf(named.read_text(encoding="utf-8", errors="replace"), out, html=True)
            res = _libreoffice(named, "pdf", tmpdir)
            shutil.move(res, out)
            return out
        finally:
            shutil.rmtree(tmpdir, ignore_errors=True)
    raise ToolError(f"Formato \"{ext}\" non supportato. Formati accettati: {', '.join(SUPPORTED_INPUT)}")


def _text_to_pdf(text: str, out: Path, html: bool) -> Path:
    import html as htmlmod

    body = text if html else "<pre style='font-family: monospace; font-size: 9pt; white-space: pre-wrap'>" \
        + htmlmod.escape(text) + "</pre>"
    story = pymupdf.Story(html=body)
    writer = pymupdf.DocumentWriter(str(out))
    mediabox = pymupdf.paper_rect("a4")
    where = mediabox + (50, 50, -50, -50)
    more = True
    while more:
        dev = writer.begin_page(mediabox)
        more, _ = story.place(where)
        story.draw(dev)
        writer.end_page()
    writer.close()
    return out


@tool("images_to_pdf", multi=True, needs_doc=False)
def images_to_pdf(ctx: Ctx) -> Result:
    """Unisce più file (immagini o documenti) in un unico PDF, una pagina per immagine."""
    ids = ctx.params.get("files") or []
    if not ids:
        raise ToolError("Scegli almeno un file.")
    fit = ctx.params.get("page_size", "image")  # image | A4
    out = pymupdf.open()
    for aid in ids:
        it = store.get(aid)
        ext = Path(it.name).suffix.lower()
        # immagini (es. foto): direttamente in memoria, senza file intermedi
        src = _image_pdf(it.path, ext) if ext in IMAGE_EXT else pymupdf.open(any_to_pdf(it.path, it.name))
        if fit == "A4":
            for i in range(src.page_count):
                r = src[i].rect
                a4 = pymupdf.paper_rect("a4-l" if r.width > r.height else "a4")
                np_ = out.new_page(width=a4.width, height=a4.height)
                np_.show_pdf_page(a4 + (20, 20, -20, -20), src, i)
        else:
            out.insert_pdf(src)
        src.close()
    path = ctx.out()
    out.save(path, garbage=1, deflate=True)
    name = str(ctx.params.get("name") or "Documento.pdf")
    if not name.lower().endswith(".pdf"):
        name += ".pdf"
    return Result(pdf=path, name=name, message=f"Creato PDF da {len(ids)} file.")


# ------------------------------------------------------------------ PDF -> altri formati


def _pdf_to_xlsx(src: Path, out: Path, pages: list[int]) -> int:
    import pdfplumber
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    n_tables = 0
    with pdfplumber.open(str(src)) as pdf:
        for pno in pages:
            page = pdf.pages[pno]
            tables = page.extract_tables()
            if not tables:
                tables = page.extract_tables({"vertical_strategy": "text", "horizontal_strategy": "text"})
                tables = [t for t in tables if t and len(t) > 1 and max(len(r) for r in t) > 1]
            for ti, table in enumerate(tables):
                n_tables += 1
                ws = wb.create_sheet(f"Pag{pno + 1}" + (f"-T{ti + 1}" if len(tables) > 1 else ""))
                for r, row in enumerate(table, 1):
                    for c, val in enumerate(row, 1):
                        v = (val or "").strip()
                        num_v = _to_number(v)
                        cell = ws.cell(row=r, column=c, value=num_v if num_v is not None else v)
                        cell.alignment = Alignment(wrap_text=True, vertical="top")
                        if r == 1:
                            cell.font = Font(bold=True)
                            cell.fill = PatternFill("solid", fgColor="E8EEF7")
                for c in range(1, max((len(r) for r in table), default=0) + 1):
                    width = max((len(str(table[r][c - 1] or "")) for r in range(len(table))
                                 if c - 1 < len(table[r])), default=8)
                    ws.column_dimensions[get_column_letter(c)].width = min(max(width + 2, 8), 60)
                ws.freeze_panes = "A2"
    if n_tables == 0:
        # nessuna tabella: testo riga per riga
        doc = pymupdf.open(src)
        ws = wb.create_sheet("Testo")
        r = 1
        for pno in pages:
            for line in doc[pno].get_text("text").splitlines():
                ws.cell(row=r, column=1, value=pno + 1)
                ws.cell(row=r, column=2, value=line)
                r += 1
        doc.close()
    wb.save(out)
    return n_tables


def _to_number(v: str):
    s = v.replace("€", "").replace("%", "").replace(" ", "").strip()
    if not s or any(ch.isalpha() for ch in s):
        return None
    # formato italiano 1.234,56
    if "," in s and s.rfind(",") > s.rfind("."):
        s = s.replace(".", "").replace(",", ".")
    else:
        s = s.replace(",", "")
    try:
        f = float(s)
        return int(f) if f.is_integer() and "." not in s else f
    except ValueError:
        return None


def _pdf_to_pptx(src: Path, out: Path, pages: list[int], dpi: int, notes: bool) -> None:
    from pptx import Presentation
    from pptx.util import Emu

    doc = pymupdf.open(src)
    first = doc[pages[0]].rect
    prs = Presentation()
    prs.slide_width = Emu(int(first.width * 12700))
    prs.slide_height = Emu(int(first.height * 12700))
    blank = prs.slide_layouts[6]
    for pno in pages:
        pg = doc[pno]
        slide = prs.slides.add_slide(blank)
        pix = pg.get_pixmap(dpi=dpi)
        img = io.BytesIO(pix.tobytes("png"))
        slide.shapes.add_picture(img, 0, 0, width=prs.slide_width, height=prs.slide_height)
        if notes:
            slide.notes_slide.notes_text_frame.text = pg.get_text("text")[:5000]
    prs.save(out)
    doc.close()


@tool("convert_from_pdf")
def convert_from_pdf(ctx: Ctx) -> Result:
    fmt = (ctx.params.get("format") or "docx").lower()
    engine = ctx.params.get("engine", "auto")
    src = ctx.item.path
    doc = ctx.open()
    if doc.is_encrypted:
        # lavora su una copia in chiaro
        tmp = ctx.out()
        doc.save(tmp, encryption=pymupdf.PDF_ENCRYPT_NONE)
        src = tmp
    pages = parse_pages(ctx.params.get("pages"), doc.page_count)
    stem = Path(ctx.item.name).stem
    doc.close()
    out = ctx.out("." + fmt if fmt not in ("png", "jpg") else ".zip")
    if fmt in ("docx", "doc", "rtf", "odt"):
        word_fmt = {"docx": 16, "doc": 0, "rtf": 6, "odt": 23}[fmt]
        if engine in ("auto", "office") and ms_office_available():
            try:
                _word_pdf_to(src, out, word_fmt)
                return Result(output=out, output_name=f"{stem}.{fmt}", message="Convertito con Microsoft Word.")
            except Exception:
                if engine == "office":
                    raise ToolError("Conversione con Microsoft Word non riuscita.")
                out.unlink(missing_ok=True)  # si riprova con il motore integrato (pdf2docx)
        if engine == "libreoffice":
            res = _libreoffice(Path(src), fmt, out.parent)
            shutil.move(res, out)
            return Result(output=out, output_name=f"{stem}.{fmt}")
        from pdf2docx import Converter

        docx_path = out if fmt == "docx" else ctx.out(".docx")
        cv = Converter(str(src))
        try:
            cv.convert(str(docx_path), pages=pages)
        finally:
            cv.close()
        if fmt != "docx":
            if not soffice_cmd():
                raise ToolError(f"Per il formato .{fmt} serve LibreOffice o Microsoft Word.")
            tmpdir = Path(tempfile.mkdtemp(dir=store.root))
            named = tmpdir / f"{stem}.docx"
            shutil.move(docx_path, named)
            res = _libreoffice(named, fmt, tmpdir)
            shutil.move(res, out)
            shutil.rmtree(tmpdir, ignore_errors=True)
        return Result(output=out, output_name=f"{stem}.{fmt}", message="Conversione completata.")
    if fmt in ("xlsx", "ods", "csv"):
        xlsx = out if fmt == "xlsx" else ctx.out(".xlsx")
        n = _pdf_to_xlsx(Path(src), xlsx, pages)
        if fmt == "csv":
            from openpyxl import load_workbook
            import csv

            wb = load_workbook(xlsx)
            with open(out, "w", newline="", encoding="utf-8-sig") as f:
                w = csv.writer(f, delimiter=";")
                for ws in wb.worksheets:
                    for row in ws.iter_rows(values_only=True):
                        w.writerow(["" if v is None else v for v in row])
                    w.writerow([])
        elif fmt == "ods":
            tmpdir = Path(tempfile.mkdtemp(dir=store.root))
            named = tmpdir / f"{stem}.xlsx"
            shutil.move(xlsx, named)
            res = _libreoffice(named, "ods", tmpdir)
            shutil.move(res, out)
            shutil.rmtree(tmpdir, ignore_errors=True)
        return Result(output=out, output_name=f"{stem}.{fmt}",
                      message=f"Estratte {n} tabelle." if n else "Nessuna tabella trovata: esportato il testo.")
    if fmt in ("pptx", "odp"):
        pptx = out if fmt == "pptx" else ctx.out(".pptx")
        _pdf_to_pptx(Path(src), pptx, pages, int(num(ctx.params.get("dpi"), 200)),
                     boolean(ctx.params.get("notes"), True))
        if fmt == "odp":
            tmpdir = Path(tempfile.mkdtemp(dir=store.root))
            named = tmpdir / f"{stem}.pptx"
            shutil.move(pptx, named)
            res = _libreoffice(named, "odp", tmpdir)
            shutil.move(res, out)
            shutil.rmtree(tmpdir, ignore_errors=True)
        return Result(output=out, output_name=f"{stem}.{fmt}", message="Presentazione creata.")
    if fmt in ("txt", "html"):
        d = pymupdf.open(src)
        if fmt == "txt":
            text = "\n\f".join(d[p].get_text("text") for p in pages)
            out.write_text(text, encoding="utf-8")
        else:
            parts = [d[p].get_text("xhtml") for p in pages]
            out.write_text("<!doctype html><meta charset='utf-8'><title>" + stem + "</title>\n"
                           + "\n<hr>\n".join(parts), encoding="utf-8")
        d.close()
        return Result(output=out, output_name=f"{stem}.{fmt}")
    if fmt in ("png", "jpg"):
        dpi = int(num(ctx.params.get("dpi"), 150))
        d = pymupdf.open(src)
        with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
            for p in pages:
                pix = d[p].get_pixmap(dpi=dpi)
                data = pix.tobytes("png") if fmt == "png" else pix.tobytes("jpeg", jpg_quality=90)
                z.writestr(f"{stem}_pag{p + 1:03d}.{fmt}", data)
        d.close()
        if len(pages) == 1:
            with zipfile.ZipFile(out) as z:
                name = z.namelist()[0]
                single = ctx.out("." + fmt)
                single.write_bytes(z.read(name))
            return Result(output=single, output_name=name)
        return Result(output=out, output_name=f"{stem}_immagini.zip")
    if fmt == "pdfa":
        gs = ghostscript_cmd()
        if not gs:
            raise ToolError("Per PDF/A serve Ghostscript (incluso nell'installer).")
        out = ctx.out(".pdf")
        level = str(ctx.params.get("pdfa_level", "2"))
        args = [gs, f"-dPDFA={level}", "-dBATCH", "-dNOPAUSE", "-dQUIET", "-sDEVICE=pdfwrite",
                "-sColorConversionStrategy=RGB", "-dPDFACompatibilityPolicy=1",
                f"-sOutputFile={out}", str(src)]
        try:
            subprocess.run(args, check=True, timeout=600, capture_output=True)
        except Exception as ex:
            raise ToolError(f"Conversione PDF/A non riuscita: {ex}")
        return Result(pdf=out, name=f"{stem} (PDF-A).pdf", message=f"Convertito in PDF/A-{level}b.")
    raise ToolError(f"Formato di destinazione non supportato: {fmt}")


@tool("file_to_pdf", needs_doc=False)
def file_to_pdf(ctx: Ctx) -> Result:
    it = ctx.asset("file")
    p = any_to_pdf(it.path, it.name, ctx.params.get("engine", "auto"))
    return Result(pdf=p, name=Path(it.name).stem + ".pdf", message="Convertito in PDF.")



